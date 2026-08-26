import codecs
import hmac
from pathlib import Path
from typing import Any

import fitz
from docx import Document

from app.services.evidence_storage_service import (
    EvidenceStorageError,
    calculate_file_sha256,
    resolve_storage_path,
)

MAX_EXTRACTED_TEXT_CHARS = 250_000
MAX_PDF_PAGES = 250


class EvidenceExtractionError(Exception):
    """Base error for evidence text extraction."""


class EvidenceExtractionIntegrityError(EvidenceExtractionError):
    """Raised when the stored evidence fails integrity checks."""


class EvidenceExtractionLimitError(EvidenceExtractionError):
    """Raised when evidence exceeds safe processing limits."""


class EvidenceExtractionUnsupportedError(EvidenceExtractionError):
    """Raised for unsupported evidence types."""


class EvidenceExtractionReadError(EvidenceExtractionError):
    """Raised when the evidence cannot be parsed."""


def normalize_extracted_text(
    text: str,
) -> str:
    """
    Normalize extracted evidence text without
    materially rewriting user-supplied content.
    """

    text = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    # Remove NUL characters.
    text = text.replace(
        "\x00",
        "",
    )

    lines = []

    for line in text.splitlines():
        cleaned = line.strip()

        if cleaned:
            lines.append(cleaned)

        elif lines and lines[-1] != "":
            lines.append("")

    while lines and lines[-1] == "":
        lines.pop()

    return "\n".join(lines)


def enforce_text_limit(
    text: str,
) -> None:
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise EvidenceExtractionLimitError(
            "Extracted evidence text exceeds " "the safe processing limit."
        )


def verify_evidence_integrity(
    evidence: dict[str, Any],
) -> Path:
    storage_path = evidence.get("storage_path")

    if not storage_path:
        raise EvidenceExtractionIntegrityError("Evidence storage path is unavailable.")

    try:
        path = resolve_storage_path(storage_path)

    except EvidenceStorageError as exc:
        raise EvidenceExtractionIntegrityError(
            "Evidence storage path is invalid."
        ) from exc

    try:
        if not path.exists() or not path.is_file():
            raise EvidenceExtractionIntegrityError(
                "Stored evidence file is unavailable."
            )

        if path.is_symlink():
            raise EvidenceExtractionIntegrityError(
                "Symbolic-link evidence files " "are not accepted."
            )

        expected_size = evidence.get("size_bytes")

        if expected_size is not None:
            actual_size = path.stat().st_size

            if actual_size != expected_size:
                raise EvidenceExtractionIntegrityError(
                    "Stored evidence failed its " "size integrity check."
                )

    except EvidenceExtractionError:
        raise

    except OSError as exc:
        raise EvidenceExtractionIntegrityError(
            "Evidence file metadata could not " "be verified."
        ) from exc

    expected_hash = str(
        evidence.get(
            "sha256",
            "",
        )
    ).lower()

    if not expected_hash:
        raise EvidenceExtractionIntegrityError(
            "Evidence integrity hash is unavailable."
        )

    try:
        actual_hash = calculate_file_sha256(storage_path)

    except (
        EvidenceStorageError,
        OSError,
    ) as exc:
        raise EvidenceExtractionIntegrityError(
            "Evidence integrity could not be verified."
        ) from exc

    if not hmac.compare_digest(
        actual_hash.lower(),
        expected_hash,
    ):
        raise EvidenceExtractionIntegrityError(
            "Stored evidence failed its SHA-256 integrity check."
        )

    return path


def extract_txt(
    path: Path,
) -> dict[str, Any]:
    decoder = codecs.getincrementaldecoder("utf-8-sig")(errors="strict")

    parts: list[str] = []
    character_count = 0

    try:
        with path.open("rb") as file:
            while True:
                chunk = file.read(64 * 1024)

                if not chunk:
                    break

                decoded = decoder.decode(
                    chunk,
                    final=False,
                )

                character_count += len(decoded)

                if character_count > MAX_EXTRACTED_TEXT_CHARS:
                    raise EvidenceExtractionLimitError(
                        "Extracted evidence text exceeds " "the safe processing limit."
                    )

                parts.append(decoded)

        final_text = decoder.decode(
            b"",
            final=True,
        )

        if final_text:
            parts.append(final_text)

    except UnicodeDecodeError as exc:
        raise EvidenceExtractionReadError(
            "Text evidence must contain valid UTF-8 text."
        ) from exc

    except OSError as exc:
        raise EvidenceExtractionReadError("Text evidence could not be read.") from exc

    text = normalize_extracted_text("".join(parts))

    enforce_text_limit(text)

    return {
        "text": text,
        "method": "utf8_text",
        "page_count": None,
    }


def extract_pdf(
    path: Path,
) -> dict[str, Any]:
    document = None

    try:
        document = fitz.open(path)

        if document.needs_pass:
            raise EvidenceExtractionReadError(
                "Encrypted PDF evidence cannot be processed."
            )

        if document.page_count > MAX_PDF_PAGES:
            raise EvidenceExtractionLimitError(
                "PDF evidence contains too many pages " "for safe processing."
            )

        parts: list[str] = []
        total_characters = 0

        for page_number in range(document.page_count):
            page = document.load_page(page_number)

            page_text = normalize_extracted_text(page.get_text("text"))

            if not page_text:
                continue

            block = f"[Page {page_number + 1}]\n" f"{page_text}"

            total_characters += len(block)

            if total_characters > MAX_EXTRACTED_TEXT_CHARS:
                raise EvidenceExtractionLimitError(
                    "Extracted evidence text exceeds " "the safe processing limit."
                )

            parts.append(block)

        text = "\n\n".join(parts)

        return {
            "text": text,
            "method": "pymupdf",
            "page_count": document.page_count,
        }

    except EvidenceExtractionError:
        raise

    except Exception as exc:
        raise EvidenceExtractionReadError("PDF evidence could not be parsed.") from exc

    finally:
        if document is not None:
            try:
                document.close()

            except Exception:
                # Cleanup failure must not mask the
                # extraction result or original error.
                pass


def extract_docx(
    path: Path,
) -> dict[str, Any]:
    try:
        document = Document(path)

        parts: list[str] = []
        total_characters = 0

        def add_text(
            value: str,
        ) -> None:
            nonlocal total_characters

            cleaned = normalize_extracted_text(value)

            if not cleaned:
                return

            total_characters += len(cleaned)

            if total_characters > MAX_EXTRACTED_TEXT_CHARS:
                raise EvidenceExtractionLimitError(
                    "Extracted evidence text " "exceeds the safe " "processing limit."
                )

            parts.append(cleaned)

        for paragraph in document.paragraphs:
            add_text(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                values = []

                for cell in row.cells:
                    value = normalize_extracted_text(cell.text)

                    if value:
                        values.append(value)

                if values:
                    add_text(" | ".join(values))

        text = "\n\n".join(parts)

        return {
            "text": text,
            "method": "python_docx",
            "page_count": None,
        }

    except EvidenceExtractionError:
        raise

    except Exception as exc:
        raise EvidenceExtractionReadError("DOCX evidence could not be parsed.") from exc


def build_no_text_result(
    method: str,
    reason: str,
    page_count=None,
) -> dict[str, Any]:
    return {
        "status": "no_text",
        "text": "",
        "character_count": 0,
        "method": method,
        "page_count": page_count,
        "reason": reason,
    }


def extract_evidence_text(
    evidence: dict[str, Any],
) -> dict[str, Any]:
    """
    Extract textual content from a previously
    validated and securely stored evidence file.

    This function does not persist extracted text.
    """

    path = verify_evidence_integrity(evidence)

    extension = str(
        evidence.get(
            "file_extension",
            "",
        )
    ).lower()

    if extension in {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
    }:
        return {
            "status": "requires_ocr",
            "text": "",
            "character_count": 0,
            "method": "none",
            "page_count": None,
            "reason": ("Image evidence requires OCR " "before textual analysis."),
        }

    if extension == ".txt":
        result = extract_txt(path)

    elif extension == ".pdf":
        result = extract_pdf(path)

    elif extension == ".docx":
        result = extract_docx(path)

    else:
        raise EvidenceExtractionUnsupportedError(
            "This evidence type does not support " "text extraction."
        )

    text = result["text"]

    if not text.strip():
        return build_no_text_result(
            method=result["method"],
            reason=("No machine-readable text was " "found in this evidence file."),
            page_count=result.get("page_count"),
        )

    return {
        "status": "ready",
        "text": text,
        "character_count": len(text),
        "method": result["method"],
        "page_count": result.get("page_count"),
        "reason": None,
    }
