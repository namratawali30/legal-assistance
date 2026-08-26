from io import BytesIO
from html import escape
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from app.repositories.complaint_repository import (
    get_complaint,
)


class ComplaintExportError(Exception):
    """Base exception for complaint export failures."""


class ComplaintNotFinalizedError(ComplaintExportError):
    """Raised when a non-finalized complaint is exported."""


class ComplaintExportContentError(ComplaintExportError):
    """Raised when finalized complaint content is unavailable."""


class ComplaintExportGenerationError(ComplaintExportError):
    """
    Raised when an otherwise exportable complaint
    cannot be rendered safely.
    """


def validate_exportable_complaint(
    complaint: dict[str, Any],
) -> None:
    status = complaint.get(
        "status",
        "draft",
    )

    if status != "finalized":
        raise ComplaintNotFinalizedError("Only finalized complaints can be exported.")

    generated_text = (complaint.get("generated_text") or "").strip()

    if not generated_text:
        raise ComplaintExportContentError("Finalized complaint text is not available.")


def build_export_filename(
    complaint_id,
    extension: str,
) -> str:
    return f"legal_complaint_" f"{str(complaint_id)}" f".{extension}"


def format_source_text(
    source: dict[str, Any],
) -> str:
    citation_id = source.get(
        "citation_id",
        "SOURCE",
    )

    title = source.get(
        "title",
        "Unknown legal source",
    )

    authority = source.get(
        "authority",
        "Unknown authority",
    )

    provision_type = source.get("provision_type")

    provision_number = source.get("provision_number")

    provision_title = source.get("provision_title")

    page_start = source.get("page_start")

    page_end = source.get("page_end")

    parts = [
        f"[{citation_id}]",
        title,
        f"Authority: {authority}",
    ]

    if provision_type and provision_number:
        parts.append(f"{provision_type.title()} " f"{provision_number}")

    if provision_title:
        parts.append(provision_title)

    if page_start is not None and page_end is not None:
        if page_start == page_end:
            parts.append(f"Page {page_start}")
        else:
            parts.append(f"Pages {page_start}-{page_end}")

    return " | ".join(parts)


def add_docx_body_text(
    document: Document,
    text: str,
) -> None:
    lines = text.splitlines()

    for line in lines:
        cleaned = line.strip()

        if not cleaned:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)

        paragraph.paragraph_format.line_spacing = 1.15

        run = paragraph.add_run(cleaned)

        run.font.name = "Times New Roman"

        run.font.size = Pt(12)

        if cleaned.startswith("Subject:") or cleaned in {
            "Complainant Details",
            "Respondent Details",
            "Facts of the Complaint",
            "Relevant Legal Basis",
            "Relief / Action Requested",
        }:
            run.bold = True


def build_docx_export(
    complaint: dict[str, Any],
) -> bytes:
    validate_exportable_complaint(complaint)

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"

    normal_style.font.size = Pt(12)

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        complaint.get(
            "title",
            "Legal Complaint",
        )
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    document.add_paragraph()

    generated_text = complaint["generated_text"]

    add_docx_body_text(
        document,
        generated_text,
    )

    sources = complaint.get(
        "sources",
        [],
    )

    if sources:
        document.add_page_break()

        heading = document.add_paragraph()

        heading_run = heading.add_run("Authoritative Legal Sources")

        heading_run.bold = True
        heading_run.font.name = "Times New Roman"
        heading_run.font.size = Pt(14)

        for source in sources:
            paragraph = document.add_paragraph(style="List Bullet")

            run = paragraph.add_run(format_source_text(source))

            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

            official_url = source.get("landing_page") or source.get("pdf_url")

            if official_url:
                url_paragraph = document.add_paragraph()

                url_run = url_paragraph.add_run(f"Official source: " f"{official_url}")

                url_run.font.name = "Times New Roman"
                url_run.font.size = Pt(9)

    document.add_paragraph()

    note = document.add_paragraph()

    note_run = note.add_run(
        "Export note: This document was prepared "
        "with AI-assisted drafting for general "
        "informational purposes and should be "
        "reviewed before formal submission."
    )

    note_run.italic = True
    note_run.font.name = "Times New Roman"
    note_run.font.size = Pt(9)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def build_pdf_export(
    complaint: dict[str, Any],
) -> bytes:
    validate_exportable_complaint(complaint)

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=complaint.get(
            "title",
            "Legal Complaint",
        ),
        author="AI Legal Assistance",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ComplaintTitle",
        parent=styles["Title"],
        fontName="Times-Bold",
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=14,
    )

    body_style = ParagraphStyle(
        "ComplaintBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=11,
        leading=16,
        spaceAfter=7,
    )

    source_heading_style = ParagraphStyle(
        "SourceHeading",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=13,
        leading=16,
        spaceAfter=10,
    )

    source_style = ParagraphStyle(
        "SourceBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=9,
        leading=12,
        spaceAfter=8,
    )

    note_style = ParagraphStyle(
        "ExportNote",
        parent=styles["BodyText"],
        fontName="Times-Italic",
        fontSize=8,
        leading=11,
        spaceBefore=12,
    )

    story = []

    story.append(
        Paragraph(
            escape(
                complaint.get(
                    "title",
                    "Legal Complaint",
                )
            ),
            title_style,
        )
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    generated_text = complaint["generated_text"]

    for line in generated_text.splitlines():
        cleaned = line.strip()

        if not cleaned:
            story.append(
                Spacer(
                    1,
                    3 * mm,
                )
            )
            continue

        story.append(
            Paragraph(
                escape(cleaned),
                body_style,
            )
        )

    sources = complaint.get(
        "sources",
        [],
    )

    if sources:
        story.append(PageBreak())

        story.append(
            Paragraph(
                "Authoritative Legal Sources",
                source_heading_style,
            )
        )

        for source in sources:
            source_text = format_source_text(source)

            story.append(
                Paragraph(
                    escape(source_text),
                    source_style,
                )
            )

            official_url = source.get("landing_page") or source.get("pdf_url")

            if official_url:
                story.append(
                    Paragraph(
                        escape("Official source: " + official_url),
                        source_style,
                    )
                )

    story.append(
        Paragraph(
            (
                "Export note: This document was "
                "prepared with AI-assisted drafting "
                "for general informational purposes "
                "and should be reviewed before "
                "formal submission."
            ),
            note_style,
        )
    )

    document.build(story)

    output.seek(0)

    return output.getvalue()


async def export_complaint_docx(
    complaint_id,
    user_id,
) -> tuple[bytes, str] | None:
    try:
        complaint = await get_complaint(
            complaint_id=complaint_id,
            user_id=user_id,
        )

    except Exception as exc:
        raise ComplaintExportGenerationError(
            "Complaint export is temporarily " "unavailable."
        ) from exc

    if not complaint:
        return None

    try:
        content = build_docx_export(complaint)

    except ComplaintExportError:
        raise

    except Exception as exc:
        raise ComplaintExportGenerationError(
            "DOCX complaint export could " "not be generated."
        ) from exc

    filename = build_export_filename(
        complaint["_id"],
        "docx",
    )

    return (
        content,
        filename,
    )


async def export_complaint_pdf(
    complaint_id,
    user_id,
) -> tuple[bytes, str] | None:
    try:
        complaint = await get_complaint(
            complaint_id=complaint_id,
            user_id=user_id,
        )

    except Exception as exc:
        raise ComplaintExportGenerationError(
            "Complaint export is temporarily " "unavailable."
        ) from exc

    if not complaint:
        return None

    try:
        content = build_pdf_export(complaint)

    except ComplaintExportError:
        raise

    except Exception as exc:
        raise ComplaintExportGenerationError(
            "PDF complaint export could " "not be generated."
        ) from exc

    filename = build_export_filename(
        complaint["_id"],
        "pdf",
    )

    return (
        content,
        filename,
    )
