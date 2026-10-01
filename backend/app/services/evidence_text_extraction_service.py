import codecs
import hmac
import logging
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from PIL import Image

from app.services.evidence_storage_service import (
    EvidenceStorageError,
    calculate_file_sha256,
    resolve_storage_path,
)

logger = logging.getLogger(__name__)

MAX_EXTRACTED_TEXT_CHARS = 250_000
MAX_PDF_PAGES = 250


class EvidenceExtractionError(Exception):
    """Base error for evidence text extraction."""


class EvidenceExtractionIntegrityError(EvidenceExtractionError):
    """Raised when the stored evidence fails integrity checks."""


class EvidenceExtractionLimitError(EvidenceExtractionError):
    """Raised when evidence exceeds safe processing limits."""


class EvidenceExtractionUnsupportedError(EvidenceExtractionError):
    """Raised for unsupported evidence types."""


class EvidenceExtractionReadError(EvidenceExtractionError):
    """Raised when the evidence cannot be parsed."""


def normalize_extracted_text(text: str) -> str:
    """
    Normalize extracted evidence text without materially rewriting user content.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")

    lines = []
    for line in text.splitlines():
        cleaned = line.strip()
        if cleaned:
            lines.append(cleaned)
        elif lines and lines[-1] != "":
            lines.append("")

    while lines and lines[-1] == "":
        lines.pop()

    return "\n".join(lines)


def enforce_text_limit(text: str) -> None:
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise EvidenceExtractionLimitError(
            "Extracted evidence text exceeds the safe processing limit."
        )


def verify_evidence_integrity(evidence: dict[str, Any]) -> Path:
    scan_status = evidence.get("scan_status")
    if scan_status in ("INFECTED", "SCAN_FAILED"):
        raise EvidenceExtractionIntegrityError("Evidence failed security scan and cannot be processed.")

    storage_path = evidence.get("storage_path")
    if not storage_path:
        raise EvidenceExtractionIntegrityError("Evidence storage path is unavailable.")

    try:
        path = resolve_storage_path(storage_path)
    except EvidenceStorageError as exc:
        raise EvidenceExtractionIntegrityError("Evidence storage path is invalid.") from exc

    try:
        if not path.exists() or not path.is_file():
            raise EvidenceExtractionIntegrityError("Stored evidence file is unavailable.")

        if path.is_symlink():
            raise EvidenceExtractionIntegrityError("Symbolic-link evidence files are not accepted.")

        expected_size = evidence.get("size_bytes")
        if expected_size is not None:
            actual_size = path.stat().st_size
            if actual_size != expected_size:
                raise EvidenceExtractionIntegrityError("Stored evidence failed its size integrity check.")

    except EvidenceExtractionError:
        raise
    except OSError as exc:
        raise EvidenceExtractionIntegrityError("Evidence file metadata could not be verified.") from exc

    expected_hash = str(evidence.get("sha256", "")).lower()
    if not expected_hash:
        raise EvidenceExtractionIntegrityError("Evidence integrity hash is unavailable.")

    try:
        actual_hash = calculate_file_sha256(storage_path)
    except (EvidenceStorageError, OSError) as exc:
        raise EvidenceExtractionIntegrityError("Evidence integrity could not be verified.") from exc

    if not hmac.compare_digest(actual_hash.lower(), expected_hash):
        raise EvidenceExtractionIntegrityError("Stored evidence failed its SHA-256 integrity check.")

    return path


def run_pytesseract_ocr(image: Image.Image) -> str:
    try:
        import pytesseract
        text = pytesseract.image_to_string(image)
        return normalize_extracted_text(text)
    except Exception as exc:
        logger.warning(f"PyTesseract OCR error: {exc}")
        return ""


def extract_image_ocr(path: Path) -> dict[str, Any]:
    try:
        with Image.open(path) as img:
            text = run_pytesseract_ocr(img)
            return {
                "text": text,
                "method": "pytesseract_ocr",
                "page_count": 1,
            }
    except Exception as exc:
        raise EvidenceExtractionReadError("Image evidence could not be processed for OCR.") from exc


def extract_txt(path: Path) -> dict[str, Any]:
    decoder = codecs.getincrementaldecoder("utf-8-sig")(errors="strict")
    parts: list[str] = []
    character_count = 0

    try:
        with path.open("rb") as file:
            while True:
                chunk = file.read(64 * 1024)
                if not chunk:
                    break
                decoded = decoder.decode(chunk, final=False)
                character_count += len(decoded)
                if character_count > MAX_EXTRACTED_TEXT_CHARS:
                    raise EvidenceExtractionLimitError("Extracted evidence text exceeds the safe processing limit.")
                parts.append(decoded)

        final_text = decoder.decode(b"", final=True)
        if final_text:
            parts.append(final_text)

    except UnicodeDecodeError as exc:
        raise EvidenceExtractionReadError("Text evidence must contain valid UTF-8 text.") from exc
    except OSError as exc:
        raise EvidenceExtractionReadError("Text evidence could not be read.") from exc

    text = normalize_extracted_text("".join(parts))
    enforce_text_limit(text)
    return {
        "text": text,
        "method": "utf8_text",
        "page_count": None,
    }


def extract_pdf(path: Path) -> dict[str, Any]:
    document = None
    used_ocr = False

    try:
        document = fitz.open(path)

        if document.needs_pass:
            raise EvidenceExtractionReadError("Encrypted PDF evidence cannot be processed.")

        if document.page_count > MAX_PDF_PAGES:
            raise EvidenceExtractionLimitError("PDF evidence contains too many pages for safe processing.")

        parts: list[str] = []
        total_characters = 0

        for page_number in range(document.page_count):
            page = document.load_page(page_number)
            page_text = normalize_extracted_text(page.get_text("text"))

            if not page_text:
                try:
                    pix = page.get_pixmap(dpi=150)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = run_pytesseract_ocr(img)
                    if ocr_text:
                        page_text = ocr_text
                        used_ocr = True
                except Exception as exc:
                    logger.warning(f"PDF page {page_number+1} OCR fallback failed: {exc}")

            if not page_text:
                continue

            block = f"[Page {page_number + 1}]\n{page_text}"
            total_characters += len(block)
            if total_characters > MAX_EXTRACTED_TEXT_CHARS:
                raise EvidenceExtractionLimitError("Extracted evidence text exceeds the safe processing limit.")

            parts.append(block)

        text = "\n\n".join(parts)
        method = "pymupdf_ocr_fallback" if used_ocr else "pymupdf"

        return {
            "text": text,
            "method": method,
            "page_count": document.page_count,
        }

    except EvidenceExtractionError:
        raise
    except Exception as exc:
        raise EvidenceExtractionReadError("PDF evidence could not be parsed.") from exc
    finally:
        if document is not None:
            try:
                document.close()
            except Exception:
                pass


def extract_docx(path: Path) -> dict[str, Any]:
    try:
        document = Document(path)
        parts: list[str] = []
        total_characters = 0

        def add_text(value: str) -> None:
            nonlocal total_characters
            cleaned = normalize_extracted_text(value)
            if not cleaned:
                return
            total_characters += len(cleaned)
            if total_characters > MAX_EXTRACTED_TEXT_CHARS:
                raise EvidenceExtractionLimitError("Extracted evidence text exceeds the safe processing limit.")
            parts.append(cleaned)

        for paragraph in document.paragraphs:
            add_text(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                values = [normalize_extracted_text(cell.text) for cell in row.cells if normalize_extracted_text(cell.text)]
                if values:
                    add_text(" | ".join(values))

        text = "\n\n".join(parts)
        return {
            "text": text,
            "method": "python_docx",
            "page_count": None,
        }
    except EvidenceExtractionError:
        raise
    except Exception as exc:
        raise EvidenceExtractionReadError("DOCX evidence could not be parsed.") from exc


def build_no_text_result(method: str, reason: str, page_count=None) -> dict[str, Any]:
    return {
        "status": "no_text",
        "text": "",
        "character_count": 0,
        "method": method,
        "page_count": page_count,
        "reason": reason,
    }


def extract_evidence_text(evidence: dict[str, Any], perform_ocr: bool = False) -> dict[str, Any]:
    """
    Extract textual content from a previously validated and securely stored evidence file.
    By default, images return 'requires_ocr' status until background OCR processing is run.
    """
    path = verify_evidence_integrity(evidence)
    extension = str(evidence.get("file_extension", "")).lower()

    if extension in {".jpg", ".jpeg", ".png", ".webp"}:
        if not perform_ocr:
            return {
                "status": "requires_ocr",
                "text": "",
                "character_count": 0,
                "method": "none",
                "page_count": None,
                "reason": "Image evidence requires OCR before textual analysis.",
            }
        result = extract_image_ocr(path)
    elif extension == ".txt":
        result = extract_txt(path)
    elif extension == ".pdf":
        result = extract_pdf(path)
    elif extension == ".docx":
        result = extract_docx(path)
    elif extension in {".mp3", ".wav", ".m4a"}:
        from app.services.transcription_service import transcribe_audio
        tx_res = transcribe_audio(path)
        if tx_res.get("status") == "no_text":
            return build_no_text_result("audio_transcription", "No speech detected in audio recording.")
        if tx_res.get("status") == "failed":
            raise EvidenceExtractionReadError("Audio transcription failed.")
        result = {
            "text": tx_res.get("transcript") or "",
            "method": "audio_transcription",
            "page_count": None,
        }
    elif extension in {".mp4", ".webm"}:
        from app.services.transcription_service import extract_audio_from_video, transcribe_audio
        temp_wav = path.parent / f".tmp_{path.stem}.wav"
        try:
            extracted = extract_audio_from_video(path, temp_wav)
            if not extracted or not temp_wav.exists():
                return build_no_text_result("video_audio_extraction", "No audio track or speech extracted from video.")
            tx_res = transcribe_audio(temp_wav)
            if tx_res.get("status") == "no_text":
                return build_no_text_result("video_audio_transcription", "No speech detected in extracted video audio.")
            if tx_res.get("status") == "failed":
                raise EvidenceExtractionReadError("Video audio transcription failed.")
            result = {
                "text": tx_res.get("transcript") or "",
                "method": "video_audio_transcription",
                "page_count": None,
            }
        finally:
            if temp_wav.exists():
                try:
                    temp_wav.unlink()
                except OSError:
                    pass
    else:
        raise EvidenceExtractionUnsupportedError("This evidence type does not support text extraction.")

    text = result["text"]

    if not text.strip():
        return build_no_text_result(
            method=result["method"],
            reason="No machine-readable text was found in this evidence file.",
            page_count=result.get("page_count"),
        )

    return {
        "status": "ready",
        "text": text,
        "character_count": len(text),
        "method": result["method"],
        "page_count": result.get("page_count"),
        "reason": None,
    }
