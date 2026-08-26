import hashlib
from pathlib import Path

import fitz
from docx import Document

from app.services import (
    evidence_storage_service,
)
from app.services.evidence_text_extraction_service import (
    EvidenceExtractionIntegrityError,
    extract_evidence_text,
)


def configure_storage(
    monkeypatch,
    tmp_path,
) -> Path:
    upload_root = (
        tmp_path
        / "uploads"
    )

    monkeypatch.setattr(
        evidence_storage_service.settings,
        "upload_dir",
        str(upload_root),
    )

    return upload_root


def build_evidence_record(
    path: Path,
    upload_root: Path,
    extension: str,
    media_type: str,
):
    content = path.read_bytes()

    relative_path = (
        path.relative_to(
            upload_root
        ).as_posix()
    )

    return {
        "storage_path":
            relative_path,

        "file_extension":
            extension,

        "media_type":
            media_type,

        "size_bytes":
            len(
                content
            ),

        "sha256":
            hashlib.sha256(
                content
            ).hexdigest(),
    }


def create_user_directory(
    upload_root: Path,
) -> Path:
    directory = (
        upload_root
        / "evidence"
        / "0123456789abcdef01234567"
    )

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    return directory


def test_txt_evidence_extraction(
    monkeypatch,
    tmp_path,
):
    upload_root = configure_storage(
        monkeypatch,
        tmp_path,
    )

    directory = create_user_directory(
        upload_root
    )

    path = (
        directory
        / "test.txt"
    )

    path.write_text(
        (
            "Invoice number ABC-123\n"
            "Amount paid: INR 25000"
        ),
        encoding="utf-8",
    )

    evidence = build_evidence_record(
        path=path,
        upload_root=upload_root,
        extension=".txt",
        media_type="text/plain",
    )

    result = extract_evidence_text(
        evidence
    )

    assert result[
        "status"
    ] == "ready"

    assert (
        "Invoice number ABC-123"
        in result[
            "text"
        ]
    )

    assert (
        "INR 25000"
        in result[
            "text"
        ]
    )

    assert result[
        "method"
    ] == "utf8_text"

    assert result[
        "character_count"
    ] > 0


def test_pdf_evidence_extraction(
    monkeypatch,
    tmp_path,
):
    upload_root = configure_storage(
        monkeypatch,
        tmp_path,
    )

    directory = create_user_directory(
        upload_root
    )

    path = (
        directory
        / "test.pdf"
    )

    document = fitz.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "Receipt number PDF-123",
    )

    document.save(
        path
    )

    document.close()

    evidence = build_evidence_record(
        path=path,
        upload_root=upload_root,
        extension=".pdf",
        media_type="application/pdf",
    )

    result = extract_evidence_text(
        evidence
    )

    assert result[
        "status"
    ] == "ready"

    assert (
        "Receipt number PDF-123"
        in result[
            "text"
        ]
    )

    assert (
        "[Page 1]"
        in result[
            "text"
        ]
    )

    assert result[
        "page_count"
    ] == 1

    assert result[
        "method"
    ] == "pymupdf"


def test_docx_evidence_extraction(
    monkeypatch,
    tmp_path,
):
    upload_root = configure_storage(
        monkeypatch,
        tmp_path,
    )

    directory = create_user_directory(
        upload_root
    )

    path = (
        directory
        / "test.docx"
    )

    document = Document()

    document.add_paragraph(
        "Employment appointment letter"
    )

    document.add_paragraph(
        "Monthly salary: INR 35000"
    )

    document.save(
        path
    )

    evidence = build_evidence_record(
        path=path,
        upload_root=upload_root,
        extension=".docx",
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )

    result = extract_evidence_text(
        evidence
    )

    assert result[
        "status"
    ] == "ready"

    assert (
        "Employment appointment letter"
        in result[
            "text"
        ]
    )

    assert (
        "Monthly salary: INR 35000"
        in result[
            "text"
        ]
    )

    assert result[
        "method"
    ] == "python_docx"


def test_image_evidence_requires_ocr(
    monkeypatch,
    tmp_path,
):
    upload_root = configure_storage(
        monkeypatch,
        tmp_path,
    )

    directory = create_user_directory(
        upload_root
    )

    path = (
        directory
        / "image.png"
    )

    # Extraction does not parse image contents;
    # integrity verification still requires a stable
    # stored artifact.
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        b"test-image-content"
    )

    evidence = build_evidence_record(
        path=path,
        upload_root=upload_root,
        extension=".png",
        media_type="image/png",
    )

    result = extract_evidence_text(
        evidence
    )

    assert result[
        "status"
    ] == "requires_ocr"

    assert result[
        "text"
    ] == ""

    assert result[
        "character_count"
    ] == 0


def test_tampered_evidence_is_not_extracted(
    monkeypatch,
    tmp_path,
):
    upload_root = configure_storage(
        monkeypatch,
        tmp_path,
    )

    directory = create_user_directory(
        upload_root
    )

    path = (
        directory
        / "tampered.txt"
    )

    path.write_text(
        "Original evidence.",
        encoding="utf-8",
    )

    evidence = build_evidence_record(
        path=path,
        upload_root=upload_root,
        extension=".txt",
        media_type="text/plain",
    )

    # Modify the stored artifact after its metadata
    # hash was calculated.
    path.write_text(
        "Modified evidence.",
        encoding="utf-8",
    )

    try:
        extract_evidence_text(
            evidence
        )

    except EvidenceExtractionIntegrityError:
        pass

    else:
        raise AssertionError(
            "Tampered evidence was unexpectedly extracted."
        )